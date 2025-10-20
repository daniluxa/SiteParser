import requests
from bs4 import BeautifulSoup as bs
import requests
import openpyxl
import time
import sys
import selenium
from selenium import webdriver
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.common.by import By


def GetDetailNumber(index, file_path='zzap/zzap.xlsx'):
    """
    Возвращает номер детали по индексу из Excel файла
    """
    try:
        # Загружаем workbook и активный лист
        wb = openpyxl.load_workbook(file_path)
        ws = wb.active
        
        # Проверяем, что индекс в пределах данных
        if index < 0 or index >= ws.max_row:
            return f"Ошибка: индекс {index} вне диапазона данных (0-{ws.max_row-1})"
        
        # Возвращаем номер детали из второго столбца (столбец B)
        return ws.cell(row=index+1, column=2).value
    
    except FileNotFoundError:
        return f"Ошибка: файл {file_path} не найден"
    except Exception as e:
        return f"Ошибка при чтении файла: {e}"

def filter_existing_companies_advanced(part_number, file_path, company_names, prices, phones, case_sensitive=False):
    try:
        wb = openpyxl.load_workbook(file_path)
        ws = wb.active
        
        part_numbers_col = 1
        companies_col = 2
        
        # Получаем список компаний из таблицы
        companies_in_table = []
        found_part = False
        
        for row in range(1, ws.max_row + 1):
            current_part = ws.cell(row=row, column=part_numbers_col).value
            
            if str(current_part) == str(part_number) and not found_part:
                found_part = True
                company = ws.cell(row=row, column=companies_col).value
                if company and str(company).strip():
                    companies_in_table.append(str(company).strip())
            
            elif found_part:
                if current_part is None or str(current_part).strip() == "":
                    company = ws.cell(row=row, column=companies_col).value
                    if company and str(company).strip():
                        companies_in_table.append(str(company).strip())
                else:
                    break
        
        if not found_part:
            print(f"Номер детали '{part_number}' не найден")
            return company_names, prices, phones
        
        # Нормализуем регистр если нужно
        if not case_sensitive:
            companies_in_table = [company.lower() for company in companies_in_table]
            company_names_normalized = [name.lower() for name in company_names]
        else:
            company_names_normalized = company_names
        
        # Фильтруем списки
        filtered_names = []
        filtered_prices = []
        filtered_phones = []
        
        removed_indices = []
        
        for i in range(len(company_names)):
            current_company = company_names_normalized[i] if not case_sensitive else company_names[i]
            table_company_to_compare = companies_in_table
            
            if not case_sensitive:
                table_company_to_compare = [c.lower() for c in companies_in_table]
            
            # Проверяем наличие компании в таблице
            if current_company not in table_company_to_compare:
                filtered_names.append(company_names[i])
                filtered_prices.append(prices[i])
                filtered_phones.append(phones[i])
            else:
                removed_indices.append(i)
        
        print(f"Для детали {part_number}:")
        print(f"Найдено в таблице: {companies_in_table}")
        print(f"Удалено компаний: {len(removed_indices)}")
        print(f"Осталось: {len(filtered_names)}")
        
        return filtered_names, filtered_prices, filtered_phones
        
    except Exception as e:
        print(f"Ошибка: {e}")
        return company_names, prices, phones

def add_new_companies_with_check(part_number, file_path, company_names, prices, phones):
    """
    Добавляет только те компании, которых еще нет в таблице
    """
    try:
        wb = openpyxl.load_workbook(file_path)
        ws = wb.active
        
        part_numbers_col = 1
        companies_col = 2
        prices_col = 3
        phones_col = 4
        
        # Собираем ВСЕ существующие компании из всего столбца 2
        all_existing_companies = set()
        
        for row in range(1, ws.max_row + 1):
            company = ws.cell(row=row, column=companies_col).value
            if company and str(company).strip():
                all_existing_companies.add(str(company).strip())
        
        print(f"Всего компаний в таблице: {len(all_existing_companies)}")
        print(f"Существующие компании: {list(all_existing_companies)}")
        
        # Находим блок для текущей детали
        start_row = None
        end_row = None
        
        for row in range(1, ws.max_row + 1):
            current_part = ws.cell(row=row, column=part_numbers_col).value
            
            if str(current_part) == str(part_number):
                start_row = row
                # Ищем конец блока
                current_check_row = row + 1
                while current_check_row <= ws.max_row:
                    next_part = ws.cell(row=current_check_row, column=part_numbers_col).value
                    if next_part is not None and str(next_part).strip() != "":
                        end_row = current_check_row
                        break
                    current_check_row += 1
                
                if end_row is None:
                    end_row = ws.max_row + 1
                break
        
        if start_row is None:
            print(f"Номер детали '{part_number}' не найден")
            return False
        
        print(f"Блок для детали {part_number}: строки {start_row}-{end_row-1}")
        
        # Фильтруем только новые компании (которых нет во всей таблице)
        new_company_names = []
        new_prices = []
        new_phones = []
        
        for i in range(len(company_names)):
            if company_names[i] not in all_existing_companies:
                new_company_names.append(company_names[i])
                new_prices.append(prices[i])
                new_phones.append(phones[i])
            else:
                print(f"Компания '{company_names[i]}' уже есть в таблице - пропускаем")
        
        if not new_company_names:
            print("Нет новых компаний для добавления")
            return True
        
        print(f"Новые компании для добавления: {new_company_names}")
        
        # Добавляем новые данные в блок текущей детали
        current_row = start_row + 1  # начинаем с первой строки после номера детали
        
        # Если текущая строка уже занята, ищем первую свободную
        while current_row < end_row:
            company_in_cell = ws.cell(row=current_row, column=companies_col).value
            if company_in_cell is None or str(company_in_cell).strip() == "":
                break
            current_row += 1
        
        # Если все строки заняты, добавляем новые
        if current_row >= end_row:
            rows_to_add = len(new_company_names)
            ws.insert_rows(current_row, rows_to_add)
            end_row += rows_to_add
        
        # Заполняем данные
        added_count = 0
        for i in range(len(new_company_names)):
            ws.cell(row=current_row, column=companies_col).value = new_company_names[i]
            ws.cell(row=current_row, column=prices_col).value = new_prices[i]
            ws.cell(row=current_row, column=phones_col).value = new_phones[i]
            
            print(f"Добавлено: {new_company_names[i]} - {new_prices[i]} - {new_phones[i]}")
            added_count += 1
            current_row += 1
        
        wb.save(file_path)
        print(f"Успешно добавлено {added_count} новых компаний")
        return True
        
    except Exception as e:
        print(f"Ошибка: {e}")
        return False

def main():

    print('Start\n')
    #Проверяем количество аргументов
    if len(sys.argv) != 3:
        print("Использование: python script.py <номер_детали> <путь_к_файлу>")
        print("Пример: python script.py 411150541 C:/Users/User/data.xlsx")
        sys.exit(1)

    part_number = sys.argv[1]
    file_path = sys.argv[2]

    #   https://www.zzap.ru/public/search.aspx#rawdata=4343060071&codes_addr=1.-1&delivery_days=0
    url = 'https://www.zzap.ru/'

    #detail_number = GetDetailNumber(0)
    detail_number = part_number
    details_name = []
    prices = []
    phone_number = []

    detail_params_list = []

    availability_of_the_next_page = True

    #new_url = f"{url}/{catalog_name[catalog_name_cnt]}/?page={page_num}"
    new_url = f"{url}/public/search.aspx#rawdata={detail_number}&codes_man=-2&codes_addr=1.-1&delivery_days=0;0.5;1"

    print("open browser\n")
    service = Service(executable_path='./geckodriver.exe')
    options = webdriver.FirefoxOptions()
    driver = webdriver.Firefox(options=options)

    driver.get(new_url)
    time.sleep(10)
    print("start parsing\n")
    while(availability_of_the_next_page):

        # Получаем HTML-код страницы
        page = driver.page_source

        soup = bs(page, "html.parser")

        details_names_tmp = soup.find_all("a", class_="f14b alink")  # извлекаем название 
        for element in details_names_tmp:
            details_name.append(element.text)

        prices_tmp = soup.find_all("span", class_="dxeBase_ZZapAqua f14b dx-nowrap")  # извлекаем цену 
        for element in prices_tmp:
            prices.append(element.text)

        phone_number_tmp = soup.find_all("span", class_="f11b")  # извлекаем телефон 
        for element in phone_number_tmp:
            phone_number.append(element.text)
        
        try:
            button = driver.find_element(By.XPATH, "//img[@class='dxWeb_pNext_ZZapAqua' and @alt='Следующая']")
            button.click()
        except:
            button = None
            availability_of_the_next_page = False


    #detail_params_list = filter_existing_companies_advanced(detail_number, 'zzap/details.xlsx', details_name, prices, phone_number)
    add_new_companies_with_check(detail_number, file_path, details_name, prices, phone_number)

    driver.quit()

if __name__ == "__main__":
    print("Запуск скрипта напрямую")
    main()
    print("Конец выполнения скрипта")